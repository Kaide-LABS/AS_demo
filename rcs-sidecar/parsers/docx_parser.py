import io
from docx import Document


async def parse(file_bytes: bytes, filename: str):
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []
    for p in doc.paragraphs:
        prefix = ""
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            level = p.style.name.replace("Heading", "").strip() or "1"
            prefix = "#" * int(level) + " "
        paragraphs.append(prefix + p.text)

    raw_text = "\n".join(paragraphs)

    tables = []
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if len(rows) > 1:
            headers = rows[0]
            for row in rows[1:]:
                tables.append(dict(zip(headers, row)))

    metadata = {"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)}
    return raw_text, tables, metadata
